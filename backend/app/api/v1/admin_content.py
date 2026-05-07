"""管理端：审核内容 CSV 导出与 Excel 批量导入（走真实识别流水线并落库）。"""
from __future__ import annotations

import csv
import io
import uuid
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from backend.app.core.deps import get_db, require_permission
from backend.app.models.orm import AuditLog, User
from backend.app.services.orchestrator import pipeline_orchestrator
from backend.app.services.review_list import list_review_items_filtered
from backend.app.services.sync_analyze_persist import persist_sync_analyze

router = APIRouter(prefix="/admin/content", tags=["admin"])
_perm = require_permission("admin:stats")

MAX_IMPORT_ROWS = 40


@router.get("/export")
def export_review_csv(
    db: Annotated[Session, Depends(get_db)],
    _: Annotated[User, Depends(_perm)],
    risk_level: str | None = Query(None),
    content_type: str | None = Query(None),
    audit_status: str | None = Query(None),
) -> StreamingResponse:
    items = list_review_items_filtered(
        db,
        risk_level=risk_level or None,
        content_type=content_type or None,
        audit_status=(audit_status or None) if audit_status != "" else None,
    )
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(
        [
            "content_id",
            "content_type",
            "risk_level",
            "llm_score",
            "audit_status",
            "username",
            "upload_time",
            "content_preview",
        ]
    )
    for row in items:
        w.writerow(
            [
                row.get("content_id"),
                row.get("content_type"),
                row.get("risk_level"),
                row.get("llm_score"),
                row.get("audit_status"),
                row.get("username"),
                row.get("upload_time"),
                (row.get("content_text") or "").replace("\n", " ")[:500],
            ]
        )
    raw = "\ufeff" + buf.getvalue()
    return StreamingResponse(
        iter([raw.encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": 'attachment; filename="lingjian_review_export.csv"'
        },
    )


@router.post("/batch-import")
def batch_import_excel(
    db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(_perm)],
    file: UploadFile = File(...),
) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            status_code=501,
            detail="服务器未安装 openpyxl，请执行 pip install openpyxl",
        )
    raw = file.file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件超过 10MB")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"无法解析 Excel：{e}") from e
    ws = wb.active
    texts: list[str] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if not row:
            continue
        cell = row[0]
        if cell is None:
            continue
        s = str(cell).strip()
        if not s:
            continue
        if idx == 0 and s in ("文本", "内容", "text", "正文"):
            continue
        texts.append(s)
        if len(texts) >= MAX_IMPORT_ROWS:
            break
    wb.close()
    if not texts:
        raise HTTPException(status_code=400, detail="未读取到有效文本行（首列应为正文，可选首行为表头）")

    ok = 0
    errors: list[str] = []
    for i, text in enumerate(texts):
        if len(text) > 8000:
            errors.append(f"第{i + 1}行超出长度上限，已跳过")
            continue
        try:
            cid = uuid.uuid4().hex
            resp = pipeline_orchestrator.analyze_text(text, cid)
            if persist_sync_analyze(db, user.id, resp):
                ok += 1
            else:
                errors.append(f"第{i + 1}行识别完成但未写入数据库，请查看后端日志（persist_sync_analyze）")
        except Exception as e:
            errors.append(f"第{i + 1}行分析失败：{e!s}"[:200])

    db.add(
        AuditLog(
            user_id=user.id,
            action="admin.content.batch_import",
            resource_type="excel",
            resource_id=file.filename or "upload.xlsx",
            detail={
                "log_type": "system",
                "imported": ok,
                "errors": errors[:30],
            },
        )
    )
    db.commit()
    return {"imported": ok, "errors": errors[:50]}
